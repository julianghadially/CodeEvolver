#!/usr/bin/env python3
"""Copyright © 2026 440 Labs LLC

Best-of-population sheet: a deterministic Excel workbook built ONLY from the
raw per-arm results files (`results/<bench>-<arm>.json`).

    python scripts/best_of_population/bop_sheet.py --results ~/ce-analysis/bop-merged/results \
        --out ~/ce-analysis/bop-merged/results/bop_sheet.xlsx [--recalc PATH/recalc.py] [--verify]

Why a script: the paper table must be reproducible from the audit's raw
observations with no hand transcription. The workbook therefore contains

  Rows        one line per (bench, arm, difficult row): the frozen winner, its
              five confirmation observations, the final program's five
              observations (blank where the final program IS the winner — that
              row's paired contribution is exactly zero), and the
              winner's selection-stage mean. These are the only raw numbers.
  Rounds      B_r / F_r / D_r per (bench, arm) as AVERAGEIFS over Rows.
  BoP         the summary table: every metric is a formula over
              Rounds; only identities (run id, |candidates|, label, historical
              final score, candidate selection means) are copied verbatim.
  Inputs      the single hard-coded constant (t critical value, from the
              results files) with its source.
  Provenance  file path + sha256 of every input, run/job ids, portfolio trees,
              testset sha256, failed-row counts, any H subsample note.
  Check       Python-recomputed metrics beside the formula cells with the
              absolute difference; MAX of those differences must be ~0.

Before writing, the script recomputes B/F/G/SE from the raw rows in Python
(mirroring `metrics.confirmation_metrics`) and refuses to write if they differ
from the metrics stored in the results files. With `--recalc` (the xlsx skill's
`recalc.py`) the workbook is recalculated by LibreOffice and `--verify` then
reads the formula results back and compares them with the Python values.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ARMS = ("freeform", "pareto_evolution")
FONT = "Arial"
NUM = "0.0000"
TOL_JSON = 1e-9      # Python recompute vs stored metrics
TOL_XLSX = 1e-6      # LibreOffice formula results vs Python


def die(msg: str) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(2)


# ---------------------------------------------------------------------------
# load + recompute
# ---------------------------------------------------------------------------

def _mean(xs: list[float]) -> float:
    return sum(xs) / len(xs)


def _sd(xs: list[float]) -> float:
    m = _mean(xs)
    return math.sqrt(sum((x - m) ** 2 for x in xs) / (len(xs) - 1))


def _se(xs: list[float]) -> float:
    return _sd(xs) / math.sqrt(len(xs))


def load_arm(path: Path, difficult_dir: Path | None) -> dict[str, Any]:
    """Parse one results file into row-level observations + Python metrics,
    and cross-check them against the metrics the report stage stored."""
    data = json.loads(path.read_text())
    bench, arm = data["bench"], data["arm"]
    reps = int(data["reps_confirm"])
    draws = list(range(int(data["reps_select"]) + 1, int(data["reps_select"]) + reps + 1))
    final_tree = data["final_tree"]
    H = sorted(int(r) for r in data["winners"])
    winner = {int(r): w["tree"] for r, w in data["winners"].items()}
    sel_mean = {int(r): float(w["mean"]) for r, w in data["winners"].items()}

    obs: dict[tuple[str, str, int, int], dict] = {}
    for rec in data["raw"]:
        key = (rec["tree"], rec["phase"], int(rec["draw"]), int(rec["row"]))
        if key in obs:
            die(f"{path.name}: duplicate raw observation {key}")
        obs[key] = rec

    def score(tree: str, phase: str, draw: int, row: int) -> float:
        rec = obs.get((tree, phase, draw, row))
        if rec is None:
            die(f"{path.name}: missing {phase} observation tree={tree[:8]} draw={draw} row={row}")
        s = rec.get("score")
        return float(s) if s is not None else 0.0   # evaluator policy: failed rows keep the recorded score (0.0)

    rows = []
    for r in H:
        w = winner[r]
        is_final = w == final_tree
        confirm = [score(w, "confirm", d, r) for d in draws]
        final = None if is_final else [score(final_tree, "final", d, r) for d in draws]
        rows.append({"row": r, "winner": w, "final_is_winner": is_final, "sel_mean": sel_mean[r],
                     "confirm": confirm, "final": final})

    B_rounds = [_mean([x["confirm"][k] for x in rows]) for k in range(reps)]
    F_rounds = [_mean([x["confirm"][k] if x["final_is_winner"] else x["final"][k] for x in rows]) for k in range(reps)]
    D_rounds = [b - f for b, f in zip(B_rounds, F_rounds)]
    m_json = data["metrics"]
    t = float(m_json["t_crit"])
    py = {
        "n_H": len(H), "B_rounds": B_rounds, "F_rounds": F_rounds, "D_rounds": D_rounds,
        "B": _mean(B_rounds), "F": _mean(F_rounds), "G": _mean(D_rounds),
        "SE_B": _se(B_rounds), "SE_F": _se(F_rounds), "SE_G": _se(D_rounds),
        "n_rows_final_is_winner": sum(1 for x in rows if x["final_is_winner"]),
        "mean_selected_selection_score": _mean([x["sel_mean"] for x in rows]),
    }
    py["sum_B"], py["sum_F"] = len(H) * py["B"], len(H) * py["F"]
    py["selection_to_confirmation_drop"] = py["mean_selected_selection_score"] - py["B"]
    py["CI95_G"] = [py["G"] - t * py["SE_G"], py["G"] + t * py["SE_G"]]

    # cross-check against what report.py stored
    for k, v in py.items():
        stored = m_json.get(k)
        if stored is None:
            die(f"{path.name}: stored metrics lack {k}")
        a = v if isinstance(v, list) else [v]
        b = stored if isinstance(stored, list) else [stored]
        if len(a) != len(b) or any(abs(float(x) - float(y)) > TOL_JSON for x, y in zip(a, b)):
            die(f"{path.name}: recomputed {k}={v} != stored {stored}")

    subsample = None
    if difficult_dir is not None:
        dpath = difficult_dir / f"{bench}.json"
        if dpath.exists():
            d = json.loads(dpath.read_text())
            if d.get("h_subsample"):
                hs = d["h_subsample"]
                subsample = f"H is a seeded uniform subsample: {hs.get('n_kept')} of {hs.get('n_full')} difficult rows (cap {hs.get('cap')}, seed {hs.get('seed')})"
            if int(d.get("n_H", len(H))) != len(H):
                die(f"{path.name}: |H| {len(H)} != difficult file n_H {d.get('n_H')}")

    portfolio_trees = [c["tree"] for c in data["portfolio"]]
    return {
        "path": path, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bench": bench, "arm": arm,
        "run_id": data["run_id"], "job_id": data["job_id"], "reported_at": data.get("reported_at"),
        "reps": reps, "draws": draws, "t_crit": t, "final_tree": final_tree, "final_forced": bool(data.get("final_forced")),
        "n_candidates": data["n_candidates_screened"], "n_unstable": len(data.get("unstable", [])),
        "label": data["difficult"]["label"], "testset_sha256": data["difficult"].get("testset_sha256"),
        "n_seed_obs": data["difficult"].get("n_seed_observations"),
        "hist_final": data["historical"]["final_test_score"], "hist_source": data["historical"]["score_source"],
        "sel_means": [data["candidate_H_means_selection"][tr] for tr in portfolio_trees],
        "portfolio_trees": portfolio_trees, "n_failed": data["n_failed_rows_by_phase"],
        "subsample": subsample, "rows": rows, "py": py,
    }


def discover(results_dirs: list[Path]) -> list[dict[str, Any]]:
    arms: dict[tuple[str, str], dict] = {}
    for d in results_dirs:
        if not d.is_dir():
            die(f"{d} is not a directory")
        difficult_dir = d.parent / "difficult"
        for p in sorted(d.glob("*.json")):
            if p.name.startswith("summary") or p.name.endswith("-rows.json"):
                continue
            a = load_arm(p, difficult_dir if difficult_dir.is_dir() else None)
            key = (a["bench"], a["arm"])
            if key in arms:
                die(f"{key} appears in both {arms[key]['path']} and {p}")
            arms[key] = a
    if not arms:
        die("no results files found")
    return [arms[k] for k in sorted(arms, key=lambda k: (k[0], ARMS.index(k[1]) if k[1] in ARMS else 9))]


# ---------------------------------------------------------------------------
# workbook
# ---------------------------------------------------------------------------

BOLD = Font(name=FONT, bold=True)
PLAIN = Font(name=FONT)
HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _set(ws, row: int, col: int, value, fmt: str | None = None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = PLAIN
    if fmt:
        cell.number_format = fmt
    return cell


def _rng(ws_name: str, col: int, first: int, last: int) -> str:
    L = get_column_letter(col)
    return f"{ws_name}!${L}${first}:${L}${last}"


def build_workbook(arms: list[dict[str, Any]], out: Path) -> dict[str, Any]:
    reps_set = {a["reps"] for a in arms}
    t_set = {a["t_crit"] for a in arms}
    if len(reps_set) != 1 or len(t_set) != 1:
        die(f"mixed reps/t_crit across inputs: reps={reps_set} t={t_set}")
    reps, t_crit = reps_set.pop(), t_set.pop()

    wb = Workbook()

    # ---- Inputs -----------------------------------------------------------
    wi = wb.active
    wi.title = "Inputs"
    _set(wi, 1, 1, "constant").font = BOLD
    _set(wi, 1, 2, "value").font = BOLD
    _set(wi, 1, 3, "source").font = BOLD
    _set(wi, 2, 1, "t_crit (95%, df = rounds - 1)")
    _set(wi, 2, 2, t_crit, "0.000").font = Font(name=FONT, color="0000FF")
    _set(wi, 2, 3, f"metrics.t_crit in every results file (spec §5: 2.776 for 4 df); rounds = {reps}")
    _set(wi, 3, 1, "confirmation rounds")
    _set(wi, 3, 2, reps).font = Font(name=FONT, color="0000FF")
    _set(wi, 3, 3, "reps_confirm in every results file")
    _set(wi, 5, 1, "Legend: blue = constant copied from the results files; black = formula; every other number lives in Rows.")
    wi.column_dimensions["A"].width = 34
    wi.column_dimensions["C"].width = 80
    T_REF = "Inputs!$B$2"

    # ---- Rows -------------------------------------------------------------
    wr = wb.create_sheet("Rows")
    hdr = ["bench", "arm", "row_id", "winner_tree", "final_is_winner", "sel_mean_winner"]
    hdr += [f"C{k + 1} (winner confirm d{a})" for k, a in enumerate(arms[0]["draws"])]
    hdr += [f"P{k + 1} (final program d{a})" for k, a in enumerate(arms[0]["draws"])]
    hdr += [f"E{k + 1} = IF(final_is_winner, C{k + 1}, P{k + 1})" for k in range(reps)]
    for c, h in enumerate(hdr, 1):
        wr.cell(row=1, column=c, value=h)
    _style_header(wr, len(hdr))
    COL_BENCH, COL_ARM, COL_ROW, COL_WIN, COL_FIW, COL_SEL = 1, 2, 3, 4, 5, 6
    COL_C0 = 7
    COL_P0 = COL_C0 + reps
    COL_E0 = COL_P0 + reps
    r = 2
    row_span: dict[tuple[str, str], tuple[int, int]] = {}
    for a in arms:
        first = r
        for x in a["rows"]:
            _set(wr, r, COL_BENCH, a["bench"])
            _set(wr, r, COL_ARM, a["arm"])
            _set(wr, r, COL_ROW, x["row"])
            _set(wr, r, COL_WIN, x["winner"][:8])
            _set(wr, r, COL_FIW, 1 if x["final_is_winner"] else 0)
            _set(wr, r, COL_SEL, x["sel_mean"], NUM)
            for k in range(reps):
                _set(wr, r, COL_C0 + k, x["confirm"][k], NUM)
                if x["final"] is not None:
                    _set(wr, r, COL_P0 + k, x["final"][k], NUM)
                cL, pL = get_column_letter(COL_C0 + k), get_column_letter(COL_P0 + k)
                fL = get_column_letter(COL_FIW)
                _set(wr, r, COL_E0 + k, f"=IF(${fL}{r}=1,{cL}{r},{pL}{r})", NUM)
            r += 1
        row_span[(a["bench"], a["arm"])] = (first, r - 1)
    ROWS_LAST = r - 1
    for c in range(1, len(hdr) + 1):
        wr.column_dimensions[get_column_letter(c)].width = 14 if c > 4 else 16

    # ---- Rounds -----------------------------------------------------------
    wo = wb.create_sheet("Rounds")
    hdr = ["bench", "arm", "n_H", "n_rows_final_is_winner", "sel_mean_of_winners"]
    hdr += [f"B_{k + 1}" for k in range(reps)] + [f"F_{k + 1}" for k in range(reps)] + [f"D_{k + 1} = B_{k + 1} - F_{k + 1}" for k in range(reps)]
    for c, h in enumerate(hdr, 1):
        wo.cell(row=1, column=c, value=h)
    _style_header(wo, len(hdr))
    RB0, RF0, RD0 = 6, 6 + reps, 6 + 2 * reps
    rounds_row: dict[tuple[str, str], int] = {}
    bench_rng = _rng("Rows", COL_BENCH, 2, ROWS_LAST)
    arm_rng = _rng("Rows", COL_ARM, 2, ROWS_LAST)
    for i, a in enumerate(arms):
        rr = i + 2
        rounds_row[(a["bench"], a["arm"])] = rr
        _set(wo, rr, 1, a["bench"])
        _set(wo, rr, 2, a["arm"])
        crit = f'{bench_rng},$A{rr},{arm_rng},$B{rr}'
        _set(wo, rr, 3, f"=COUNTIFS({crit})")
        _set(wo, rr, 4, f"=SUMIFS({_rng('Rows', COL_FIW, 2, ROWS_LAST)},{crit})")
        _set(wo, rr, 5, f"=AVERAGEIFS({_rng('Rows', COL_SEL, 2, ROWS_LAST)},{crit})", NUM)
        for k in range(reps):
            _set(wo, rr, RB0 + k, f"=AVERAGEIFS({_rng('Rows', COL_C0 + k, 2, ROWS_LAST)},{crit})", NUM)
            _set(wo, rr, RF0 + k, f"=AVERAGEIFS({_rng('Rows', COL_E0 + k, 2, ROWS_LAST)},{crit})", NUM)
            bL, fL = get_column_letter(RB0 + k), get_column_letter(RF0 + k)
            _set(wo, rr, RD0 + k, f"={bL}{rr}-{fL}{rr}", NUM)
    for c in range(1, len(hdr) + 1):
        wo.column_dimensions[get_column_letter(c)].width = 13
    wo.column_dimensions["B"].width = 18

    def rounds_range(rr: int, c0: int) -> str:
        return f"Rounds!${get_column_letter(c0)}${rr}:${get_column_letter(c0 + reps - 1)}${rr}"

    # ---- BoP --------------------------------------------------------------
    wbop = wb.create_sheet("BoP", 0)
    hdr = ["bench", "arm", "run", "cands", "|H|", "label", "F", "c1", "c2", "c3", "c4", "c5", "B", "SE(B)", "G", "SE(G)",
           "CI95(G) lo", "CI95(G) hi", "sum B", "sum F", "drop", "hist. final", "hist. source", "evo-ff B", "SE(evo-ff)",
           "n final wins", "failed rows", "final forced", "H note"]
    for c, h in enumerate(hdr, 1):
        wbop.cell(row=1, column=c, value=h)
    _style_header(wbop, len(hdr))
    C = {h: i + 1 for i, h in enumerate(hdr)}
    bop_row: dict[tuple[str, str], int] = {}
    for i, a in enumerate(arms):
        br = i + 2
        bop_row[(a["bench"], a["arm"])] = br
    for a in arms:
        br = bop_row[(a["bench"], a["arm"])]
        rr = rounds_row[(a["bench"], a["arm"])]
        _set(wbop, br, C["bench"], a["bench"])
        _set(wbop, br, C["arm"], a["arm"])
        _set(wbop, br, C["run"], a["run_id"])
        _set(wbop, br, C["cands"], a["n_candidates"])
        _set(wbop, br, C["|H|"], f"=Rounds!$C${rr}")
        _set(wbop, br, C["label"], a["label"])
        Frng, Brng, Drng = rounds_range(rr, RF0), rounds_range(rr, RB0), rounds_range(rr, RD0)
        _set(wbop, br, C["F"], f"=AVERAGE({Frng})", NUM)
        for k in range(5):
            v = a["sel_means"][k] if k < len(a["sel_means"]) else None
            _set(wbop, br, C[f"c{k + 1}"], v, NUM)
        _set(wbop, br, C["B"], f"=AVERAGE({Brng})", NUM)
        _set(wbop, br, C["SE(B)"], f"=STDEV({Brng})/SQRT(COUNT({Brng}))", NUM)
        _set(wbop, br, C["G"], f"=AVERAGE({Drng})", NUM)
        _set(wbop, br, C["SE(G)"], f"=STDEV({Drng})/SQRT(COUNT({Drng}))", NUM)
        gL, sgL = get_column_letter(C["G"]), get_column_letter(C["SE(G)"])
        _set(wbop, br, C["CI95(G) lo"], f"={gL}{br}-{T_REF}*{sgL}{br}", NUM)
        _set(wbop, br, C["CI95(G) hi"], f"={gL}{br}+{T_REF}*{sgL}{br}", NUM)
        hL, bL, fL = get_column_letter(C["|H|"]), get_column_letter(C["B"]), get_column_letter(C["F"])
        _set(wbop, br, C["sum B"], f"={hL}{br}*{bL}{br}", NUM)
        _set(wbop, br, C["sum F"], f"={hL}{br}*{fL}{br}", NUM)
        _set(wbop, br, C["drop"], f"=Rounds!$E${rr}-{bL}{br}", NUM)
        _set(wbop, br, C["hist. final"], a["hist_final"], NUM)
        _set(wbop, br, C["hist. source"], a["hist_source"])
        other = bop_row.get((a["bench"], "pareto_evolution")), bop_row.get((a["bench"], "freeform"))
        if other[0] and other[1]:
            sbL = get_column_letter(C["SE(B)"])
            _set(wbop, br, C["evo-ff B"], f"={bL}{other[0]}-{bL}{other[1]}", NUM)
            _set(wbop, br, C["SE(evo-ff)"], f"=SQRT({sbL}{other[0]}^2+{sbL}{other[1]}^2)", NUM)
        _set(wbop, br, C["n final wins"], f"=Rounds!$D${rr}")
        _set(wbop, br, C["failed rows"], sum(int(v) for v in a["n_failed"].values()))
        _set(wbop, br, C["final forced"], "yes" if a["final_forced"] else "no")
        _set(wbop, br, C["H note"], a["subsample"] or "full difficult set")
    note_row = len(arms) + 3
    notes = [
        "B = mean over confirmation rounds of each row's frozen winner; F = final program on the same rows; G = B - F (paired per round).",
        "SE = sample SD over the rounds (n-1) / sqrt(rounds); CI95 = estimate +/- t_crit * SE (Inputs!B2). evo-ff B uses SE = sqrt(SE_evo^2 + SE_ff^2).",
        "drop = mean selection-stage score of the winners - B (selection-to-confirmation drop; a noisy diagnostic, not a correction).",
        "c1..c5 = each portfolio candidate's selection-stage mean on H (copied from the results files); hist. final = the run's recorded held-out score, not an audit measurement.",
        "Every metric cell is a formula over Rounds -> Rows; the Rows sheet holds the only raw observations. Provenance lists each source file's sha256.",
    ]
    for j, n in enumerate(notes):
        _set(wbop, note_row + j, 1, n).font = Font(name=FONT, italic=True)
    widths = {"bench": 12, "arm": 17, "run": 36, "label": 15, "hist. source": 26, "H note": 60}
    for h, c in C.items():
        wbop.column_dimensions[get_column_letter(c)].width = widths.get(h, 10)

    # ---- Provenance -------------------------------------------------------
    wp = wb.create_sheet("Provenance")
    hdr = ["bench", "arm", "results file", "sha256", "run_id", "job_id", "reported_at (UTC)", "cands screened", "unstable",
           "final tree", "final forced", "portfolio trees", "|H|", "label", "testset sha256", "seed observations",
           "failed rows by phase", "H note"]
    for c, h in enumerate(hdr, 1):
        wp.cell(row=1, column=c, value=h)
    _style_header(wp, len(hdr))
    for i, a in enumerate(arms):
        pr = i + 2
        ts = dt.datetime.fromtimestamp(a["reported_at"], dt.timezone.utc).strftime("%Y-%m-%d %H:%M:%S") if a["reported_at"] else ""
        vals = [a["bench"], a["arm"], str(a["path"]), a["sha256"], a["run_id"], a["job_id"], ts, a["n_candidates"], a["n_unstable"],
                a["final_tree"][:8], "yes" if a["final_forced"] else "no", ", ".join(t[:8] for t in a["portfolio_trees"]),
                a["py"]["n_H"], a["label"], a["testset_sha256"], a["n_seed_obs"], json.dumps(a["n_failed"], sort_keys=True),
                a["subsample"] or "full difficult set"]
        for c, v in enumerate(vals, 1):
            _set(wp, pr, c, v)
    for c in range(1, len(hdr) + 1):
        wp.column_dimensions[get_column_letter(c)].width = 18
    wp.column_dimensions["C"].width = 60
    wp.column_dimensions["D"].width = 40

    # ---- Check ------------------------------------------------------------
    wc = wb.create_sheet("Check")
    metrics = ["B", "F", "G", "SE_B", "SE_G", "sum_B", "sum_F", "selection_to_confirmation_drop", "n_H", "n_rows_final_is_winner"]
    bop_col = {"B": "B", "F": "F", "G": "G", "SE_B": "SE(B)", "SE_G": "SE(G)", "sum_B": "sum B", "sum_F": "sum F",
               "selection_to_confirmation_drop": "drop", "n_H": "|H|", "n_rows_final_is_winner": "n final wins"}
    hdr = ["bench", "arm", "metric", "python (from raw)", "formula (BoP)", "abs diff"]
    for c, h in enumerate(hdr, 1):
        wc.cell(row=1, column=c, value=h)
    _style_header(wc, len(hdr))
    cr = 2
    for a in arms:
        br = bop_row[(a["bench"], a["arm"])]
        for mname in metrics:
            _set(wc, cr, 1, a["bench"])
            _set(wc, cr, 2, a["arm"])
            _set(wc, cr, 3, mname)
            _set(wc, cr, 4, a["py"][mname], NUM)
            L = get_column_letter(C[bop_col[mname]])
            _set(wc, cr, 5, f"=BoP!${L}${br}", NUM)
            _set(wc, cr, 6, f"=ABS(D{cr}-E{cr})", "0.000000000")
            cr += 1
    _set(wc, cr + 1, 3, "MAX abs diff").font = BOLD
    _set(wc, cr + 1, 6, f"=MAX(F2:F{cr - 1})", "0.000000000").font = BOLD
    wc.column_dimensions["C"].width = 32
    for L in "DEF":
        wc.column_dimensions[L].width = 18

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"bop_row": bop_row, "bop_cols": C, "check_max_cell": f"F{cr + 1}"}


# ---------------------------------------------------------------------------
# recalc + verify
# ---------------------------------------------------------------------------

def recalc(script: Path, out: Path) -> None:
    res = subprocess.run([sys.executable, str(script), str(out)], capture_output=True, text=True)
    print(res.stdout.strip())
    try:
        rep = json.loads(res.stdout.strip().splitlines()[-1])
    except Exception:  # noqa: BLE001
        die(f"recalc produced no JSON report: {res.stderr[-400:]}")
    if rep.get("status") != "success" or rep.get("total_errors", 0):
        die(f"recalc reported errors: {rep}")


def verify(arms: list[dict[str, Any]], out: Path, layout: dict[str, Any]) -> None:
    wb = load_workbook(out, data_only=True)

    def sheet(name: str):  # the `formulas` package writes sheet names upper-cased
        for ws in wb.worksheets:
            if ws.title.lower() == name.lower():
                return ws
        die(f"{out}: no sheet {name!r}")

    ws = sheet("BoP")
    C = layout["bop_cols"]
    checks = {"B": "B", "F": "F", "G": "G", "SE_B": "SE(B)", "SE_G": "SE(G)", "sum_B": "sum B", "sum_F": "sum F",
              "selection_to_confirmation_drop": "drop", "n_H": "|H|", "n_rows_final_is_winner": "n final wins"}
    worst = 0.0
    n = 0
    for a in arms:
        br = layout["bop_row"][(a["bench"], a["arm"])]
        for k, h in checks.items():
            v = ws.cell(row=br, column=C[h]).value
            if v is None:
                die(f"{a['bench']}/{a['arm']} {h}: formula has no cached value (run --recalc first)")
            d = abs(float(v) - float(a["py"][k]))
            worst = max(worst, d)
            n += 1
            if d > TOL_XLSX:
                die(f"{a['bench']}/{a['arm']} {h}: sheet {v} != python {a['py'][k]} (diff {d})")
        lo = ws.cell(row=br, column=C["CI95(G) lo"]).value
        hi = ws.cell(row=br, column=C["CI95(G) hi"]).value
        for got, exp in ((lo, a["py"]["CI95_G"][0]), (hi, a["py"]["CI95_G"][1])):
            if got is None or abs(float(got) - exp) > TOL_XLSX:
                die(f"{a['bench']}/{a['arm']} CI95(G): sheet {got} != python {exp}")
            n += 1
    mx = sheet("Check")[layout["check_max_cell"]].value
    print(f"verify: {n} formula cells match the Python recomputation (max abs diff {worst:.2e}); Check!MAX = {mx}")


# ---------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results", action="append", required=True, help="results/ directory (repeatable)")
    ap.add_argument("--out", required=True, help="output .xlsx")
    ap.add_argument("--recalc", help="path to the xlsx skill's recalc.py (LibreOffice recalculation)")
    ap.add_argument("--verify", action="store_true", help="read formula results back and compare with Python (needs cached values)")
    ap.add_argument("--verify-file", help="verify THIS recalculated copy instead of --out (e.g. one written by the `formulas` package)")
    args = ap.parse_args(argv)

    arms = discover([Path(r).expanduser() for r in args.results])
    for a in arms:
        p = a["py"]
        print(f"  {a['bench']:12s} {a['arm']:17s} |H|={p['n_H']:4d}  B={p['B']:.4f}  F={p['F']:.4f}  G={p['G']:+.4f} (SE {p['SE_G']:.4f})"
              f"  [recomputed from {len(a['rows'])} rows x {a['reps']} rounds; matches stored metrics]")
    out = Path(args.out).expanduser()
    layout = build_workbook(arms, out)
    print(f"wrote {out}  ({len(arms)} bench/arm rows)")
    if args.recalc:
        recalc(Path(args.recalc).expanduser(), out)
    if args.verify or args.recalc or args.verify_file:
        verify(arms, Path(args.verify_file).expanduser() if args.verify_file else out, layout)
    return 0


if __name__ == "__main__":
    sys.exit(main())
